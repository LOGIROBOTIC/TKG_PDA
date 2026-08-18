from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SCHEMA_VERSION = "phase_02a_v1"
SOURCE_ROOT = Path("2026- OPTİMAK -ROBOT AMR")
WORK_PLAN = SOURCE_ROOT / "AMR Varyant" / "AMR Çalışma Planı.xlsx"
TAKT_STUDY = SOURCE_ROOT / "AMR Takt time study.xlsx"
INTERACTION = SOURCE_ROOT / "AGV - BAKIM - PLC Çalışma.xlsx"
LAYOUT_WORKBOOKS = (
    SOURCE_ROOT / "AMR Varyant" / "PAH-01-A-412B.xlsx",
    SOURCE_ROOT / "AMR Varyant" / "PAH-02-A-GAR.xlsx",
    SOURCE_ROOT / "AMR Varyant" / "SH-01-A HÜCRESİ VARYANT KUTU YERLEŞİMİ.xlsx",
    SOURCE_ROOT / "AMR Varyant" / "SH-01-B HÜCRESİ VARYANT KUTU YERLEŞİMİ.xlsx",
    SOURCE_ROOT / "AMR Varyant" / "SH-01-P1--SH-01-P2-SH-02-P1 -SH-02-P2.xlsx",
)
SOURCE_FILES = (INTERACTION, TAKT_STUDY, WORK_PLAN, *LAYOUT_WORKBOOKS)

BASELINE_VARIANTS: dict[str, tuple[str, ...]] = {
    "SH-01-A": ("1", "2", "3", "4", "5"),
    "SH-01-B": ("1", "2", "3"),
    "SH-01-P1--SH-01-P2": ("1", "2"),
    "SH-02-P1 -SH-02-P2": ("1", "2"),
    "SEALER": ("1", "2", "3"),
    "PAH-01-A": ("1", "2"),
    "412B": ("1", "2", "3"),
    "PAH-02-A": ("1", "2"),
    "GAR": ("1", "2", "3"),
}

CANONICAL_PART_ALIASES = {
    "67115BN050-10-HMTR": "67115BN050-10-HAOS",
}

PART_CODE_RE = re.compile(r"^(?=.*\d)[A-Z0-9]{5,}(?:-[A-Z0-9]+)*$")
ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
DECIMAL_QUANTUM = Decimal("0.001")
TURKISH_ASCII = str.maketrans(
    {
        "ı": "i",
        "İ": "I",
        "ş": "s",
        "Ş": "S",
        "ğ": "g",
        "Ğ": "G",
        "ç": "c",
        "Ç": "C",
        "ö": "o",
        "Ö": "O",
        "ü": "u",
        "Ü": "U",
    }
)


class StructuralExtractionError(RuntimeError):
    """Raised only when generation cannot preserve referential meaning."""


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("–", "-").replace("—", "-").replace("−", "-")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def normalize_label(value: Any) -> str:
    text = normalize_text(value) or ""
    text = unicodedata.normalize("NFKD", text.translate(TURKISH_ASCII))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).strip().upper()
    return re.sub(r"\s+", " ", text)


def normalize_identifier(value: Any) -> str | None:
    text = normalize_text(value)
    if text is None:
        return None
    return text.upper().replace(" ", "")


def normalize_variant(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if number == number.to_integral_value():
            return str(int(number))
    text = normalize_text(value)
    if text is None:
        return None
    match = re.search(r"\d+", text)
    return str(int(match.group(0))) if match else text.upper()


def split_part_codes(value: Any) -> list[str]:
    if value is None:
        return []
    tokens = re.split(r"[\r\n]+", str(value))
    result: list[str] = []
    for token in tokens:
        code = normalize_identifier(token)
        if code and PART_CODE_RE.fullmatch(code) and ("-" in code or any(char.isalpha() for char in code)):
            result.append(code)
    return result


def decimal_string(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        text = normalize_text(value)
        if text is None or text.upper() in ERROR_VALUES:
            return None
        text = text.replace(",", ".")
    else:
        text = str(value)
    try:
        number = Decimal(text).quantize(DECIMAL_QUANTUM, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None
    return format(number, "f")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_ref(project_root: Path, path: Path, sheet: str, row: int, cell: str) -> dict[str, Any]:
    return {
        "file": path.relative_to(project_root).as_posix(),
        "sheet": sheet,
        "row": row,
        "cell": cell,
    }


def source_ref_key(ref: dict[str, Any]) -> tuple[Any, ...]:
    return (ref["file"], ref["sheet"], ref["row"], ref["cell"])


def stable_unique_refs(refs: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    unique = {source_ref_key(ref): ref for ref in refs}
    return [unique[key] for key in sorted(unique)]


def make_issue_id(payload: dict[str, Any]) -> str:
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:20]


def add_issue(issues: dict[str, dict[str, Any]], code: str, severity: str, message: str, **context: Any) -> None:
    payload = {"code": code, "severity": severity, "message": message, **context}
    issue_id = make_issue_id(payload)
    issues[issue_id] = {"issue_id": issue_id, "blocking": False, **payload}


def canonicalize_part(
    raw_code: str,
    issues: dict[str, dict[str, Any]],
    ref: dict[str, Any],
    machine_code: str | None = None,
    variant_code: str | None = None,
) -> str:
    canonical = CANONICAL_PART_ALIASES.get(raw_code, raw_code)
    if canonical != raw_code:
        add_issue(
            issues,
            "PART_CODE_SOURCE_DISCREPANCY",
            "WARNING",
            f"Source code {raw_code} is retained as evidence and normalized to canonical {canonical}.",
            source=ref,
            machine_code=machine_code,
            variant_code=variant_code,
            raw_part_code=raw_code,
            canonical_part_code=canonical,
        )
    return canonical


def build_manifest(project_root: Path) -> list[dict[str, Any]]:
    manifest: list[dict[str, Any]] = []
    for relative in SOURCE_FILES:
        path = project_root / relative
        if not path.is_file():
            raise StructuralExtractionError(f"Required source file is missing: {relative.as_posix()}")
        manifest.append(
            {
                "path": relative.as_posix(),
                "size_bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "role": (
                    "authoritative_textual_master_and_planning"
                    if relative == WORK_PLAN
                    else "planning_corroboration"
                    if relative == TAKT_STUDY
                    else "interaction_corroboration"
                    if relative == INTERACTION
                    else "textual_corroboration_only_no_geometry"
                ),
            }
        )
    return sorted(manifest, key=lambda item: item["path"].casefold())


def manifest_digest(manifest: list[dict[str, Any]]) -> str:
    payload = json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def find_header_row(ws: Any) -> int:
    for row_number in range(1, min(ws.max_row, 25) + 1):
        labels = {normalize_label(ws.cell(row_number, column).value) for column in range(1, ws.max_column + 1)}
        if "VARYANT" in labels and "KUTU LOTU" in labels and "PARCA" in labels:
            return row_number
    raise StructuralExtractionError(f"Could not find textual planning header in sheet {ws.title}")


def find_column(labels: dict[int, str], predicate: Any, required: bool = True) -> int | None:
    matches = [column for column, label in labels.items() if predicate(label)]
    if not matches:
        if required:
            raise StructuralExtractionError("Required planning column is missing")
        return None
    return min(matches)


def planning_columns(ws: Any, header_row: int) -> dict[str, int | None]:
    labels = {
        column: normalize_label(ws.cell(header_row, column).value)
        for column in range(1, ws.max_column + 1)
        if ws.cell(header_row, column).value is not None
    }
    main = find_column(labels, lambda label: label == "PARCA")
    related = find_column(labels, lambda label: label == "ALT PARCA", required=False) or (main + 1)
    container = find_column(labels, lambda label: label == "KASA TIPI", required=False) or (related + 1)
    rounded = find_column(
        labels,
        lambda label: "GUNLUK URETIM KASA" in label and ("YUV" in label or "YUVAR" in label),
    )
    type_hint = rounded + 1 if rounded + 1 <= ws.max_column else None
    return {
        "dimensions": find_column(labels, lambda label: "MAMUL KASA OLCU" in label),
        "main_part": main,
        "related_part": related,
        "container_code": container,
        "lot_quantity": find_column(labels, lambda label: label == "KUTU LOTU"),
        "cycle_time_seconds": find_column(labels, lambda label: label == "CEVRIM"),
        "annual_volume": find_column(labels, lambda label: label == "Y VOLUME"),
        "variant": find_column(labels, lambda label: label == "VARYANT"),
        "mix_pallet_need": find_column(labels, lambda label: "MIX PALET IHTIYACI" in label),
        "weekly_volume": find_column(labels, lambda label: "HAFTALIK VOLUME" in label),
        "daily_volume": find_column(labels, lambda label: label == "GUNLUK VOLUME"),
        "daily_production_time": find_column(labels, lambda label: "GUNLUK URETIM SURESI" in label),
        "daily_container_need": find_column(
            labels,
            lambda label: "GUNLUK URETIM KASA IHTIYACI" in label and "YUV" not in label,
        ),
        "rounded_container_need": rounded,
        "container_category_hint": type_hint,
        "agv_moves_per_day": find_column(labels, lambda label: "AGV" in label and "GUNLUK" in label),
        "area_need": find_column(
            labels,
            lambda label: "ALAN IHTIYACI" in label,
            required=False,
        ),
    }


def category_from_hint(value: Any, container_code: str | None) -> str | None:
    hint = normalize_label(value)
    if hint == "KUTU":
        return "BOX"
    if hint == "KASA":
        return "CONTAINER"
    if hint == "ARABA" or container_code == "ARABA":
        return "CART"
    return None


def merge_value(
    record: dict[str, Any],
    field: str,
    incoming: Any,
    issues: dict[str, dict[str, Any]],
    context: dict[str, Any],
    ref: dict[str, Any],
) -> None:
    if field in record.setdefault("_conflicted_fields", set()):
        return
    current = record.get(field)
    if incoming is None:
        return
    if current is None:
        record[field] = incoming
        return
    if current != incoming:
        record[field] = None
        record["_conflicted_fields"].add(field)
        add_issue(
            issues,
            "AUTHORITATIVE_DUPLICATE_VALUE_CONFLICT",
            "WARNING",
            f"Conflicting authoritative values for {field}; normalized value is NULL.",
            field=field,
            first_value=current,
            conflicting_value=incoming,
            source=ref,
            **context,
        )


def cell_value(ws: Any, row: int, column: int | None) -> Any:
    return ws.cell(row, column).value if column else None


def extract_work_plan(
    project_root: Path,
    issues: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    path = project_root / WORK_PLAN
    formula_wb = load_workbook(path, data_only=False, read_only=False)
    value_wb = load_workbook(path, data_only=True, read_only=False)

    parts: dict[str, dict[str, Any]] = {}
    groups: dict[tuple[str, str, str], dict[str, Any]] = {}
    members: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    planning: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    container_observations: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"categories": set(), "source_refs": []}
    )

    group_order_by_variant: Counter[tuple[str, str]] = Counter()
    member_order_by_group: Counter[tuple[str, str, str]] = Counter()

    def add_part(code: str, ref: dict[str, Any]) -> None:
        record = parts.setdefault(code, {"part_code": code, "source_refs": []})
        record["source_refs"].append(ref)

    try:
        for formula_ws in formula_wb.worksheets:
            if normalize_label(formula_ws.title) == "OZET":
                continue
            machine_code = formula_ws.title
            if machine_code not in BASELINE_VARIANTS:
                raise StructuralExtractionError(f"Unknown machine sheet in work plan: {machine_code}")
            value_ws = value_wb[formula_ws.title]
            header_row = find_header_row(value_ws)
            columns = planning_columns(value_ws, header_row)
            current_group_keys: list[tuple[str, str, str]] = []
            current_variant: str | None = None

            for row_number in range(header_row + 1, value_ws.max_row + 1):
                main_raw = split_part_codes(cell_value(value_ws, row_number, columns["main_part"]))
                related_raw = split_part_codes(cell_value(value_ws, row_number, columns["related_part"]))
                if not main_raw and not related_raw:
                    continue

                variant_code = normalize_variant(cell_value(value_ws, row_number, columns["variant"]))
                if variant_code is None:
                    variant_code = current_variant
                if variant_code is None or variant_code not in BASELINE_VARIANTS[machine_code]:
                    raise StructuralExtractionError(
                        f"Impossible variant reference at {WORK_PLAN.as_posix()}#{formula_ws.title}!{row_number}: "
                        f"{variant_code!r}"
                    )

                part_column = columns["main_part"] if main_raw else columns["related_part"]
                ref = source_ref(
                    project_root,
                    path,
                    formula_ws.title,
                    row_number,
                    formula_ws.cell(row_number, part_column).coordinate,
                )

                if main_raw:
                    current_variant = variant_code
                    current_group_keys = []
                    for raw_code in main_raw:
                        main_code = canonicalize_part(raw_code, issues, ref, machine_code, variant_code)
                        group_key = (machine_code, variant_code, main_code)
                        if group_key not in groups:
                            group_order_by_variant[(machine_code, variant_code)] += 1
                            groups[group_key] = {
                                "machine_code": machine_code,
                                "variant_code": variant_code,
                                "main_part_code": main_code,
                                "name": main_code,
                                "display_order": group_order_by_variant[(machine_code, variant_code)],
                                "source_refs": [],
                            }
                        groups[group_key]["source_refs"].append(ref)
                        current_group_keys.append(group_key)
                        add_part(main_code, ref)
                    row_codes = [(key, key[2], False) for key in current_group_keys]
                else:
                    if not current_group_keys or current_variant != variant_code:
                        raise StructuralExtractionError(
                            f"Related part has no valid product-group context at "
                            f"{WORK_PLAN.as_posix()}#{formula_ws.title}!{row_number}"
                        )
                    row_codes = []
                    for raw_code in related_raw:
                        related_code = canonicalize_part(raw_code, issues, ref, machine_code, variant_code)
                        add_part(related_code, ref)
                        for group_key in current_group_keys:
                            member_key = (*group_key, related_code)
                            if member_key not in members:
                                member_order_by_group[group_key] += 1
                                members[member_key] = {
                                    "machine_code": machine_code,
                                    "variant_code": variant_code,
                                    "main_part_code": group_key[2],
                                    "part_code": related_code,
                                    "relation_type": "ALT_PART",
                                    "display_order": member_order_by_group[group_key],
                                    "container_code": None,
                                    "lot_quantity": None,
                                    "source_refs": [],
                                    "_conflicted_fields": set(),
                                }
                            members[member_key]["source_refs"].append(ref)
                            row_codes.append((group_key, related_code, True))

                container_code = normalize_identifier(cell_value(value_ws, row_number, columns["container_code"]))
                category = category_from_hint(
                    cell_value(value_ws, row_number, columns["container_category_hint"]),
                    container_code,
                )
                if container_code:
                    observation = container_observations[container_code]
                    if category:
                        observation["categories"].add(category)
                    observation["source_refs"].append(ref)

                metric_values = {
                    "container_dimensions": normalize_text(
                        cell_value(value_ws, row_number, columns["dimensions"])
                    ),
                    "container_code": container_code,
                    "lot_quantity": decimal_string(
                        cell_value(value_ws, row_number, columns["lot_quantity"])
                    ),
                    "cycle_time_seconds": decimal_string(
                        cell_value(value_ws, row_number, columns["cycle_time_seconds"])
                    ),
                    "annual_volume": decimal_string(
                        cell_value(value_ws, row_number, columns["annual_volume"])
                    ),
                    "weekly_volume": decimal_string(
                        cell_value(value_ws, row_number, columns["weekly_volume"])
                    ),
                    "daily_volume": decimal_string(
                        cell_value(value_ws, row_number, columns["daily_volume"])
                    ),
                    "mix_pallet_need": decimal_string(
                        cell_value(value_ws, row_number, columns["mix_pallet_need"])
                    ),
                    "agv_moves_per_day": decimal_string(
                        cell_value(value_ws, row_number, columns["agv_moves_per_day"])
                    ),
                    "area_need": decimal_string(
                        cell_value(value_ws, row_number, columns["area_need"])
                    ),
                }

                for group_key, part_code, is_related in row_codes:
                    planning_key = (*group_key, part_code)
                    planning_record = planning.setdefault(
                        planning_key,
                        {
                            "machine_code": machine_code,
                            "variant_code": variant_code,
                            "main_part_code": group_key[2],
                            "part_code": part_code,
                            "container_dimensions": None,
                            "container_code": None,
                            "lot_quantity": None,
                            "cycle_time_seconds": None,
                            "annual_volume": None,
                            "weekly_volume": None,
                            "daily_volume": None,
                            "mix_pallet_need": None,
                            "agv_moves_per_day": None,
                            "area_need": None,
                            "source_refs": [],
                            "_conflicted_fields": set(),
                        },
                    )
                    planning_record["source_refs"].append(ref)
                    context = {
                        "machine_code": machine_code,
                        "variant_code": variant_code,
                        "main_part_code": group_key[2],
                        "part_code": part_code,
                    }
                    for field, value in metric_values.items():
                        merge_value(planning_record, field, value, issues, context, ref)

                    if is_related:
                        member = members[(*group_key, part_code)]
                        merge_value(member, "container_code", container_code, issues, context, ref)
                        merge_value(
                            member,
                            "lot_quantity",
                            metric_values["lot_quantity"],
                            issues,
                            context,
                            ref,
                        )

                    if container_code is None:
                        add_issue(
                            issues,
                            "UNKNOWN_CONTAINER_MAPPING",
                            "WARNING",
                            "Container mapping is blank in the authoritative source; no container is guessed.",
                            source=ref,
                            machine_code=machine_code,
                            variant_code=variant_code,
                            main_part_code=group_key[2],
                            part_code=part_code,
                            raw_container_code=None,
                        )
                    if metric_values["lot_quantity"] is None:
                        add_issue(
                            issues,
                            "MISSING_PLANNING_INPUT",
                            "WARNING",
                            "Kutu Lotu is blank or invalid; dependent planning metrics remain NULL.",
                            source=ref,
                            machine_code=machine_code,
                            variant_code=variant_code,
                            main_part_code=group_key[2],
                            part_code=part_code,
                            field="lot_quantity",
                        )

        return {
            "parts": parts,
            "groups": groups,
            "members": members,
            "planning": planning,
            "container_observations": container_observations,
        }
    finally:
        formula_wb.close()
        value_wb.close()


def scan_formula_errors(
    project_root: Path,
    relative_path: Path,
    issues: dict[str, dict[str, Any]],
) -> None:
    path = project_root / relative_path
    formula_wb = load_workbook(path, data_only=False, read_only=False)
    value_wb = load_workbook(path, data_only=True, read_only=False)
    try:
        for formula_ws in formula_wb.worksheets:
            value_ws = value_wb[formula_ws.title]
            for row in formula_ws.iter_rows():
                for formula_cell in row:
                    if formula_cell.data_type != "f":
                        continue
                    cached = value_ws[formula_cell.coordinate].value
                    if not (isinstance(cached, str) and cached.upper() in ERROR_VALUES):
                        continue
                    ref = source_ref(
                        project_root,
                        path,
                        formula_ws.title,
                        formula_cell.row,
                        formula_cell.coordinate,
                    )
                    add_issue(
                        issues,
                        "EXCEL_FORMULA_ERROR",
                        "WARNING",
                        "Excel cached formula result is an error; normalized metric value is NULL.",
                        source=ref,
                        formula=formula_cell.value,
                        excel_error=cached,
                    )
    finally:
        formula_wb.close()
        value_wb.close()


def extract_takt_rows(project_root: Path, issues: dict[str, dict[str, Any]]) -> dict[tuple[str, str, str, str], dict[str, Any]]:
    path = project_root / TAKT_STUDY
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        ws = wb["main"]
        current_group: dict[tuple[str, str], str] = {}
        rows: dict[tuple[str, str, str, str], dict[str, Any]] = {}
        for row_number in range(2, ws.max_row + 1):
            machine_code = normalize_text(ws.cell(row_number, 1).value)
            variant_code = normalize_variant(ws.cell(row_number, 10).value)
            main_codes = split_part_codes(ws.cell(row_number, 4).value)
            related_codes = split_part_codes(ws.cell(row_number, 5).value)
            if not machine_code or machine_code not in BASELINE_VARIANTS or not variant_code:
                continue
            if variant_code not in BASELINE_VARIANTS[machine_code]:
                raise StructuralExtractionError(
                    f"Impossible variant reference in takt study at row {row_number}: "
                    f"{machine_code}/{variant_code}"
                )
            part_column = 4 if main_codes else 5
            ref = source_ref(
                project_root,
                path,
                ws.title,
                row_number,
                ws.cell(row_number, part_column).coordinate,
            )
            if main_codes:
                raw_main = main_codes[0]
                main_code = canonicalize_part(raw_main, issues, ref, machine_code, variant_code)
                current_group[(machine_code, variant_code)] = main_code
                codes = [main_code]
            elif related_codes:
                main_code = current_group.get((machine_code, variant_code))
                if not main_code:
                    continue
                codes = [canonicalize_part(code, issues, ref, machine_code, variant_code) for code in related_codes]
            else:
                continue

            for part_code in codes:
                key = (machine_code, variant_code, main_code, part_code)
                rows[key] = {
                    "machine_code": machine_code,
                    "variant_code": variant_code,
                    "main_part_code": main_code,
                    "part_code": part_code,
                    "container_dimensions": normalize_text(ws.cell(row_number, 3).value),
                    "container_code": normalize_identifier(ws.cell(row_number, 6).value),
                    "lot_quantity": decimal_string(ws.cell(row_number, 7).value),
                    "cycle_time_seconds": decimal_string(ws.cell(row_number, 8).value),
                    "annual_volume": decimal_string(ws.cell(row_number, 9).value),
                    "mix_pallet_need": decimal_string(ws.cell(row_number, 11).value),
                    "weekly_volume": decimal_string(ws.cell(row_number, 12).value),
                    "daily_volume": decimal_string(ws.cell(row_number, 13).value),
                    "agv_moves_per_day": decimal_string(ws.cell(row_number, 18).value),
                    "source": ref,
                }
        return rows
    finally:
        wb.close()


def compare_takt_to_work_plan(
    authoritative: dict[tuple[str, str, str, str], dict[str, Any]],
    takt_rows: dict[tuple[str, str, str, str], dict[str, Any]],
    issues: dict[str, dict[str, Any]],
) -> None:
    fields = (
        "container_dimensions",
        "container_code",
        "lot_quantity",
        "cycle_time_seconds",
        "annual_volume",
        "mix_pallet_need",
        "weekly_volume",
        "daily_volume",
        "agv_moves_per_day",
    )
    for key in sorted(takt_rows):
        corroborating = takt_rows[key]
        primary = authoritative.get(key)
        context = {
            "machine_code": key[0],
            "variant_code": key[1],
            "main_part_code": key[2],
            "part_code": key[3],
        }
        if primary is None:
            add_issue(
                issues,
                "TAKT_CONTEXT_NOT_IN_WORK_PLAN",
                "INFO",
                "Takt-study row has no identical product-group context in the authoritative work plan; it is not imported independently.",
                source=corroborating["source"],
                **context,
            )
            continue
        for field in fields:
            left = primary.get(field)
            right = corroborating.get(field)
            if left is not None and right is not None and left != right:
                add_issue(
                    issues,
                    "SOURCE_VALUE_CONFLICT",
                    "WARNING",
                    f"Work-plan and takt-study values differ for {field}; work-plan precedence is retained.",
                    field=field,
                    work_plan_value=left,
                    takt_study_value=right,
                    source=corroborating["source"],
                    **context,
                )


def scan_textual_corroboration(
    project_root: Path,
    authoritative_parts: set[str],
    issues: dict[str, dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    evidence: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for relative_path in (INTERACTION, *LAYOUT_WORKBOOKS):
        path = project_root / relative_path
        wb = load_workbook(path, data_only=True, read_only=False)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        for raw_code in split_part_codes(cell.value):
                            ref = source_ref(project_root, path, ws.title, cell.row, cell.coordinate)
                            canonical = canonicalize_part(raw_code, issues, ref)
                            evidence[canonical].append(ref)
        finally:
            wb.close()

    for part_code, refs in sorted(evidence.items()):
        if part_code not in authoritative_parts:
            add_issue(
                issues,
                "CORROBORATION_ONLY_PART",
                "INFO",
                "Part appears only in a corroboration/layout source and is not imported without authoritative textual planning context.",
                part_code=part_code,
                sources=stable_unique_refs(refs),
            )
    return {code: stable_unique_refs(refs) for code, refs in sorted(evidence.items())}


def finalize_containers(
    observations: dict[str, dict[str, Any]],
    issues: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    containers: list[dict[str, Any]] = []
    for code in sorted(observations):
        categories = sorted(observations[code]["categories"])
        refs = stable_unique_refs(observations[code]["source_refs"])
        if len(categories) == 1:
            category = categories[0]
        else:
            category = "OTHER"
            add_issue(
                issues,
                "CONTAINER_CATEGORY_UNRESOLVED" if not categories else "CONTAINER_CATEGORY_CONFLICT",
                "INFO" if not categories else "WARNING",
                "Container code is preserved exactly, but its category is unresolved; category OTHER is used."
                if not categories
                else "Container code has conflicting category hints; category OTHER is used.",
                container_code=code,
                observed_categories=categories,
                sources=refs,
            )
        containers.append(
            {
                "code": code,
                "name": code,
                "category": category,
                "source_refs": refs,
            }
        )
    return containers


def source_note(record: dict[str, Any]) -> str:
    refs = stable_unique_refs(record.get("source_refs", []))
    locations = ", ".join(
        f"{ref['file']}#{ref['sheet']}!{ref['cell']}" for ref in refs
    )
    missing = [
        field
        for field in (
            "container_code",
            "lot_quantity",
            "cycle_time_seconds",
            "annual_volume",
            "weekly_volume",
            "daily_volume",
            "mix_pallet_need",
            "agv_moves_per_day",
            "area_need",
        )
        if record.get(field) is None
    ]
    notes = [f"Phase 02A source: {locations}"]
    if missing:
        notes.append("NULL fields: " + ", ".join(missing))
    if record.get("part_code") == "67115BN050-10-HAOS":
        notes.append("Canonical HAOS; isolated HMTR layout occurrence awaits customer confirmation")
    return "; ".join(notes)


def strip_internal_fields(record: dict[str, Any]) -> dict[str, Any]:
    clean = {key: value for key, value in record.items() if not key.startswith("_")}
    if "source_refs" in clean:
        clean["source_refs"] = stable_unique_refs(clean["source_refs"])
    return clean


def build_dataset(project_root: Path) -> dict[str, Any]:
    project_root = project_root.resolve()
    manifest_before = build_manifest(project_root)
    issues: dict[str, dict[str, Any]] = {}

    extracted = extract_work_plan(project_root, issues)
    scan_formula_errors(project_root, WORK_PLAN, issues)
    scan_formula_errors(project_root, TAKT_STUDY, issues)
    takt_rows = extract_takt_rows(project_root, issues)
    compare_takt_to_work_plan(extracted["planning"], takt_rows, issues)
    corroboration = scan_textual_corroboration(project_root, set(extracted["parts"]), issues)
    containers = finalize_containers(extracted["container_observations"], issues)

    parts = [strip_internal_fields(extracted["parts"][key]) for key in sorted(extracted["parts"])]
    groups = [strip_internal_fields(extracted["groups"][key]) for key in sorted(extracted["groups"])]
    members = [strip_internal_fields(extracted["members"][key]) for key in sorted(extracted["members"])]
    planning: list[dict[str, Any]] = []
    for key in sorted(extracted["planning"]):
        record = strip_internal_fields(extracted["planning"][key])
        record["source_note"] = source_note(record)
        planning.append(record)

    manifest_after = build_manifest(project_root)
    if manifest_before != manifest_after:
        raise StructuralExtractionError("Customer source files changed during extraction")

    issue_list = [issues[key] for key in sorted(issues)]
    issue_codes = Counter(issue["code"] for issue in issue_list)
    issue_severities = Counter(issue["severity"] for issue in issue_list)
    unresolved_container_codes = sorted(
        {
            issue.get("container_code") or "(blank)"
            for issue in issue_list
            if issue["code"] in {
                "UNKNOWN_CONTAINER_MAPPING",
                "CONTAINER_CATEGORY_UNRESOLVED",
                "CONTAINER_CATEGORY_CONFLICT",
            }
        }
    )

    counts = {
        "source_files": len(manifest_before),
        "unique_parts": len(parts),
        "container_types": len(containers),
        "product_groups": len(groups),
        "product_group_parts": len(members),
        "planning_rows": len(planning),
        "data_issues": len(issue_list),
        "formula_error_cells": issue_codes.get("EXCEL_FORMULA_ERROR", 0),
    }

    return {
        "schema_version": SCHEMA_VERSION,
        "source_digest": manifest_digest(manifest_before),
        "source_files_unchanged": True,
        "manifest": manifest_before,
        "precedence": [
            "AMR Çalışma Planı.xlsx is authoritative for machine/variant/product-group/planning rows.",
            "AMR Takt time study.xlsx corroborates field values and reports conflicts without overriding the work plan.",
            "AGV - BAKIM - PLC Çalışma.xlsx and layout workbooks provide textual corroboration only.",
            "No colored-cell coordinates, footprints, merged-area geometry or layout slot positions are extracted.",
        ],
        "counts": counts,
        "issue_counts_by_code": dict(sorted(issue_codes.items())),
        "issue_counts_by_severity": dict(sorted(issue_severities.items())),
        "unresolved_container_codes": unresolved_container_codes,
        "parts": parts,
        "container_types": containers,
        "product_groups": groups,
        "product_group_parts": members,
        "planning_metrics": planning,
        "corroboration_part_evidence": corroboration,
        "data_issues": issue_list,
    }


def sql_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "true" if value else "false"
    return "'" + str(value).replace("'", "''") + "'"


def sql_number(value: str | None) -> str:
    return value if value is not None else "NULL"


def values_clause(rows: Iterable[Iterable[str]]) -> str:
    return ",\n".join("    (" + ", ".join(row) + ")" for row in rows)


def generate_seed_sql(dataset: dict[str, Any]) -> str:
    parts_values = values_clause(
        (sql_literal(part["part_code"]), "NULL") for part in dataset["parts"]
    )
    container_values = values_clause(
        (
            sql_literal(item["code"]),
            sql_literal(item["name"]),
            sql_literal(item["category"]),
        )
        for item in dataset["container_types"]
    )
    group_values = values_clause(
        (
            sql_literal(item["machine_code"]),
            sql_literal(item["variant_code"]),
            sql_literal(item["main_part_code"]),
            sql_literal(item["name"]),
            str(item["display_order"]),
        )
        for item in dataset["product_groups"]
    )
    member_values = values_clause(
        (
            sql_literal(item["machine_code"]),
            sql_literal(item["variant_code"]),
            sql_literal(item["main_part_code"]),
            sql_literal(item["part_code"]),
            sql_literal(item["relation_type"]),
            sql_literal(item["container_code"]),
            sql_number(item["lot_quantity"]),
            str(item["display_order"]),
        )
        for item in dataset["product_group_parts"]
    )
    planning_values = values_clause(
        (
            sql_literal(item["machine_code"]),
            sql_literal(item["variant_code"]),
            sql_literal(item["main_part_code"]),
            sql_literal(item["part_code"]),
            sql_literal(item["container_code"]),
            sql_literal(item["container_dimensions"]),
            sql_number(item["lot_quantity"]),
            sql_number(item["cycle_time_seconds"]),
            sql_number(item["annual_volume"]),
            sql_number(item["weekly_volume"]),
            sql_number(item["daily_volume"]),
            sql_number(item["mix_pallet_need"]),
            sql_number(item["agv_moves_per_day"]),
            sql_number(item["area_need"]),
            sql_literal(item["source_note"]),
        )
        for item in dataset["planning_metrics"]
    )

    return f"""-- TKG PDA Phase 02A generated textual customer master data
-- schema_version: {dataset['schema_version']}
-- source_digest: {dataset['source_digest']}
-- deterministic output; do not edit by hand

BEGIN;

INSERT INTO master.parts (part_code, description)
VALUES
{parts_values}
ON CONFLICT (part_code) DO NOTHING;

INSERT INTO master.container_types (code, name, category)
VALUES
{container_values}
ON CONFLICT (code) DO UPDATE
SET name = EXCLUDED.name,
    category = EXCLUDED.category,
    updated_at = now()
WHERE master.container_types.name IS DISTINCT FROM EXCLUDED.name
   OR master.container_types.category IS DISTINCT FROM EXCLUDED.category;

WITH source(machine_code, variant_code, main_part_code, name, display_order) AS (
VALUES
{group_values}
)
INSERT INTO master.product_groups (variant_id, main_part_id, name, display_order)
SELECT v.id, p.id, source.name, source.display_order
FROM source
JOIN master.machines AS m ON m.code = source.machine_code
JOIN master.variants AS v ON v.machine_id = m.id AND v.code = source.variant_code
JOIN master.parts AS p ON p.part_code = source.main_part_code
ON CONFLICT (variant_id, main_part_id) DO UPDATE
SET name = EXCLUDED.name,
    display_order = EXCLUDED.display_order,
    updated_at = now()
WHERE master.product_groups.name IS DISTINCT FROM EXCLUDED.name
   OR master.product_groups.display_order IS DISTINCT FROM EXCLUDED.display_order;

WITH source(
    machine_code, variant_code, main_part_code, part_code,
    relation_type, container_code, lot_quantity, display_order
) AS (
VALUES
{member_values}
)
INSERT INTO master.product_group_parts (
    product_group_id, part_id, container_type_id,
    relation_type, lot_quantity, display_order
)
SELECT pg.id, related.id, ct.id,
       source.relation_type, source.lot_quantity, source.display_order
FROM source
JOIN master.machines AS m ON m.code = source.machine_code
JOIN master.variants AS v ON v.machine_id = m.id AND v.code = source.variant_code
JOIN master.parts AS main_part ON main_part.part_code = source.main_part_code
JOIN master.product_groups AS pg
  ON pg.variant_id = v.id AND pg.main_part_id = main_part.id
JOIN master.parts AS related ON related.part_code = source.part_code
LEFT JOIN master.container_types AS ct ON ct.code = source.container_code
ON CONFLICT (product_group_id, part_id) DO UPDATE
SET container_type_id = EXCLUDED.container_type_id,
    relation_type = EXCLUDED.relation_type,
    lot_quantity = EXCLUDED.lot_quantity,
    display_order = EXCLUDED.display_order,
    updated_at = now()
WHERE master.product_group_parts.container_type_id IS DISTINCT FROM EXCLUDED.container_type_id
   OR master.product_group_parts.relation_type IS DISTINCT FROM EXCLUDED.relation_type
   OR master.product_group_parts.lot_quantity IS DISTINCT FROM EXCLUDED.lot_quantity
   OR master.product_group_parts.display_order IS DISTINCT FROM EXCLUDED.display_order;

WITH source(
    machine_code, variant_code, main_part_code, part_code,
    container_code, container_dimensions, lot_quantity, cycle_time_seconds,
    annual_volume, weekly_volume, daily_volume, mix_pallet_need,
    agv_moves_per_day, area_need, source_note
) AS (
VALUES
{planning_values}
)
INSERT INTO master.part_planning_metrics (
    variant_id, product_group_id, part_id, container_type_id,
    container_dimensions, lot_quantity, cycle_time_seconds,
    annual_volume, weekly_volume, daily_volume, mix_pallet_need,
    agv_moves_per_day, area_need, source_note
)
SELECT v.id, pg.id, part.id, ct.id,
       source.container_dimensions, source.lot_quantity, source.cycle_time_seconds,
       source.annual_volume, source.weekly_volume, source.daily_volume,
       source.mix_pallet_need, source.agv_moves_per_day, source.area_need,
       source.source_note
FROM source
JOIN master.machines AS m ON m.code = source.machine_code
JOIN master.variants AS v ON v.machine_id = m.id AND v.code = source.variant_code
JOIN master.parts AS main_part ON main_part.part_code = source.main_part_code
JOIN master.product_groups AS pg
  ON pg.variant_id = v.id AND pg.main_part_id = main_part.id
JOIN master.parts AS part ON part.part_code = source.part_code
LEFT JOIN master.container_types AS ct ON ct.code = source.container_code
ON CONFLICT (variant_id, product_group_id, part_id)
WHERE product_group_id IS NOT NULL
DO UPDATE SET
    container_type_id = EXCLUDED.container_type_id,
    container_dimensions = EXCLUDED.container_dimensions,
    lot_quantity = EXCLUDED.lot_quantity,
    cycle_time_seconds = EXCLUDED.cycle_time_seconds,
    annual_volume = EXCLUDED.annual_volume,
    weekly_volume = EXCLUDED.weekly_volume,
    daily_volume = EXCLUDED.daily_volume,
    mix_pallet_need = EXCLUDED.mix_pallet_need,
    agv_moves_per_day = EXCLUDED.agv_moves_per_day,
    area_need = EXCLUDED.area_need,
    source_note = EXCLUDED.source_note,
    updated_at = now()
WHERE master.part_planning_metrics.container_type_id IS DISTINCT FROM EXCLUDED.container_type_id
   OR master.part_planning_metrics.container_dimensions IS DISTINCT FROM EXCLUDED.container_dimensions
   OR master.part_planning_metrics.lot_quantity IS DISTINCT FROM EXCLUDED.lot_quantity
   OR master.part_planning_metrics.cycle_time_seconds IS DISTINCT FROM EXCLUDED.cycle_time_seconds
   OR master.part_planning_metrics.annual_volume IS DISTINCT FROM EXCLUDED.annual_volume
   OR master.part_planning_metrics.weekly_volume IS DISTINCT FROM EXCLUDED.weekly_volume
   OR master.part_planning_metrics.daily_volume IS DISTINCT FROM EXCLUDED.daily_volume
   OR master.part_planning_metrics.mix_pallet_need IS DISTINCT FROM EXCLUDED.mix_pallet_need
   OR master.part_planning_metrics.agv_moves_per_day IS DISTINCT FROM EXCLUDED.agv_moves_per_day
   OR master.part_planning_metrics.area_need IS DISTINCT FROM EXCLUDED.area_need
   OR master.part_planning_metrics.source_note IS DISTINCT FROM EXCLUDED.source_note;

COMMIT;
"""


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            serialized = {
                key: json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
                if isinstance(value, (dict, list, tuple, set))
                else value
                for key, value in row.items()
            }
            writer.writerow(serialized)


def write_outputs(dataset: dict[str, Any], output_dir: Path, seed_path: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    write_json(output_dir / "normalized_customer_data.json", dataset)
    write_json(output_dir / "source_manifest.json", dataset["manifest"])
    write_json(output_dir / "summary.json", {
        "schema_version": dataset["schema_version"],
        "source_digest": dataset["source_digest"],
        "source_files_unchanged": dataset["source_files_unchanged"],
        "counts": dataset["counts"],
        "issue_counts_by_code": dataset["issue_counts_by_code"],
        "issue_counts_by_severity": dataset["issue_counts_by_severity"],
        "unresolved_container_codes": dataset["unresolved_container_codes"],
    })
    write_json(output_dir / "data_issues.json", dataset["data_issues"])

    write_csv(output_dir / "parts.csv", dataset["parts"], ["part_code"])
    write_csv(
        output_dir / "container_types.csv",
        dataset["container_types"],
        ["code", "name", "category"],
    )
    write_csv(
        output_dir / "product_groups.csv",
        dataset["product_groups"],
        ["machine_code", "variant_code", "main_part_code", "name", "display_order"],
    )
    write_csv(
        output_dir / "product_group_parts.csv",
        dataset["product_group_parts"],
        [
            "machine_code",
            "variant_code",
            "main_part_code",
            "part_code",
            "relation_type",
            "container_code",
            "lot_quantity",
            "display_order",
        ],
    )
    write_csv(
        output_dir / "planning_metrics.csv",
        dataset["planning_metrics"],
        [
            "machine_code",
            "variant_code",
            "main_part_code",
            "part_code",
            "container_code",
            "container_dimensions",
            "lot_quantity",
            "cycle_time_seconds",
            "annual_volume",
            "weekly_volume",
            "daily_volume",
            "mix_pallet_need",
            "agv_moves_per_day",
            "area_need",
            "source_note",
        ],
    )
    write_csv(
        output_dir / "data_issues.csv",
        dataset["data_issues"],
        [
            "issue_id",
            "severity",
            "code",
            "blocking",
            "message",
            "machine_code",
            "variant_code",
            "main_part_code",
            "part_code",
            "field",
            "container_code",
            "raw_container_code",
            "work_plan_value",
            "takt_study_value",
            "source",
            "sources",
        ],
    )

    seed_path.parent.mkdir(parents=True, exist_ok=True)
    seed_path.write_text(generate_seed_sql(dataset), encoding="utf-8", newline="\n")


def run_pipeline(project_root: Path, output_dir: Path, seed_path: Path) -> dict[str, Any]:
    dataset = build_dataset(project_root)
    write_outputs(dataset, output_dir, seed_path)
    return dataset


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate TKG PDA Phase 02A normalized data and SQL seed")
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(__file__).resolve().parents[2],
    )
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--seed-sql", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    project_root = args.project_root.resolve()
    output_dir = (args.output_dir or project_root / "data" / "generated" / "phase_02a").resolve()
    seed_path = (
        args.seed_sql
        or project_root / "db" / "seeds" / "generated" / "002a_customer_textual_master_data.sql"
    ).resolve()
    dataset = run_pipeline(project_root, output_dir, seed_path)
    print(json.dumps({
        "source_digest": dataset["source_digest"],
        "counts": dataset["counts"],
        "unresolved_container_codes": dataset["unresolved_container_codes"],
        "output_dir": str(output_dir),
        "seed_sql": str(seed_path),
    }, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
